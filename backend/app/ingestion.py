import logging, uuid
import fitz  # PyMuPDF
from qdrant_client.models import PointStruct
from .core import SessionLocal
from .models import Source, Chunk
from .providers import LocalEmbeddingProvider, ensure_collection, upsert

log = logging.getLogger(__name__)

def chunk_text(text: str, size=900, overlap=150):
    text = " ".join(text.split())
    i = 0
    while i < len(text):
        yield text[i:i + size]
        i += size - overlap

def ingest_pdf(source_id: str):
    db = SessionLocal()
    src = db.get(Source, source_id)
    try:
        src.status = "PROCESSING"; db.commit()
        emb = LocalEmbeddingProvider()
        doc = fitz.open(src.path)
        src.pages = len(doc)
        rows = []
        for pno, page in enumerate(doc, start=1):
            for t in chunk_text(page.get_text()):
                if len(t.strip()) > 40:
                    rows.append(Chunk(source_id=src.id, course_id=src.course_id, page=pno, text=t))
        if rows:
            vecs = emb.embed([r.text for r in rows])
            ensure_collection(len(vecs[0]))
            upsert([PointStruct(id=str(uuid.uuid4()), vector=v, payload={
                "course_id": r.course_id, "source_id": r.source_id, "source_name": src.name,
                "page": r.page, "text": r.text}) for r, v in zip(rows, vecs)])
            db.add_all(rows)
        src.status = "COMPLETED"
    except Exception as e:
        log.exception("ingest failed")
        src.status = "FAILED"; src.error = "PDF processing failed."
    db.commit(); db.close()

import subprocess, json, os
import cv2
import pytesseract

def extract_video_metadata(filepath: str) -> dict:
    cmd = ["ffprobe", "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams", filepath]
    try:
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        if result.returncode != 0: return {}
        info = json.loads(result.stdout)
    except Exception:
        return {}
    
    video_stream = next((s for s in info.get("streams", []) if s.get("codec_type") == "video"), None)
    audio_stream = next((s for s in info.get("streams", []) if s.get("codec_type") == "audio"), None)
    
    duration = float(info.get("format", {}).get("duration", 0))
    if duration == 0 and video_stream:
        duration = float(video_stream.get("duration", 0))
        
    width = video_stream.get("width") if video_stream else None
    height = video_stream.get("height") if video_stream else None
    fps = 0
    if video_stream and video_stream.get("r_frame_rate"):
        try: fps = eval(video_stream["r_frame_rate"])
        except: pass
    
    return {
        "duration": duration,
        "width": width,
        "height": height,
        "fps": fps,
        "video_codec": video_stream.get("codec_name") if video_stream else None,
        "audio_codec": audio_stream.get("codec_name") if audio_stream else None,
        "audio_sample_rate": audio_stream.get("sample_rate") if audio_stream else None,
        "audio_streams_count": sum(1 for s in info.get("streams", []) if s.get("codec_type") == "audio"),
        "size": int(info.get("format", {}).get("size", 0))
    }

def ingest_local_video(source_id: str):
    db = SessionLocal()
    src = db.get(Source, source_id)
    try:
        src.status = "PROCESSING"
        db.commit()
        
        # 1. Metadata
        meta = extract_video_metadata(src.path)
        src.metadata_json = meta
        db.commit()
        
        rows = []
        
        # 2. Extract Audio
        audio_path = src.path + ".wav"
        subprocess.run(["ffmpeg", "-y", "-i", src.path, "-vn", "-acodec", "pcm_s16le", "-ar", "16000", "-ac", "1", audio_path], 
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        
        has_audio = os.path.exists(audio_path) and os.path.getsize(audio_path) > 1000
        
        if has_audio:
            from faster_whisper import WhisperModel
            model = WhisperModel("base", device="cpu", compute_type="int8")
            segments, info = model.transcribe(audio_path, beam_size=5)
            
            for segment in segments:
                text = segment.text.strip()
                if len(text) > 5:
                    rows.append(Chunk(
                        source_id=src.id, 
                        course_id=src.course_id,
                        timestamp_start=segment.start,
                        timestamp_end=segment.end,
                        text=text
                    ))
            os.remove(audio_path)
            
        # 3. Extract Frames & OCR
        cap = cv2.VideoCapture(src.path)
        fps = cap.get(cv2.CAP_PROP_FPS) or 30
        frame_interval = max(int(fps * 30), 1) # every 30 seconds
        
        frame_count = 0
        success, image = cap.read()
        while success:
            if frame_count % frame_interval == 0:
                timestamp = frame_count / fps
                rgb_image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
                ocr_text = pytesseract.image_to_string(rgb_image).strip()
                
                if len(ocr_text) > 20: 
                    frame_filename = f"{src.id}_f{frame_count}.jpg"
                    frame_path = os.path.join("./uploads", frame_filename)
                    cv2.imwrite(frame_path, image)
                    rows.append(Chunk(
                        source_id=src.id, 
                        course_id=src.course_id,
                        timestamp_start=timestamp,
                        timestamp_end=timestamp + 5,
                        text=ocr_text,
                        ocr_text=ocr_text,
                        frame_path=frame_path
                    ))
            success, image = cap.read()
            frame_count += 1
        cap.release()
        
        # 4. Embeddings & DB
        if rows:
            emb = LocalEmbeddingProvider()
            vecs = emb.embed([r.text for r in rows])
            ensure_collection(len(vecs[0]))
            upsert([PointStruct(id=str(uuid.uuid4()), vector=v, payload={
                "course_id": r.course_id, "source_id": r.source_id, "source_name": src.name,
                "timestamp_start": r.timestamp_start, "timestamp_end": r.timestamp_end,
                "text": r.text}) for r, v in zip(rows, vecs)])
            db.add_all(rows)
            
        src.status = "COMPLETED"
    except Exception as e:
        log.exception("local video ingest failed")
        src.status = "FAILED"
        src.error = str(e)
    finally:
        db.commit()
        db.close()

def ingest_youtube(source_id: str):
    db = SessionLocal()
    src = db.get(Source, source_id)
    try:
        src.status = "PROCESSING"; db.commit()
        import yt_dlp, os
        from faster_whisper import WhisperModel
        
        ydl_opts = {
            'format': 'bestaudio',
            'outtmpl': f'./uploads/{src.id}.%(ext)s',
        }
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(src.path, download=True)
            ext = info.get('ext', 'm4a')
            
        audio_path = f'./uploads/{src.id}.{ext}'
        
        model = WhisperModel("base", device="cpu", compute_type="int8")
        segments, info = model.transcribe(audio_path, beam_size=5)

        rows = []
        for segment in segments:
            text = segment.text.strip()
            if len(text) > 5:
                rows.append(Chunk(
                    source_id=src.id, 
                    course_id=src.course_id,
                    timestamp_start=segment.start,
                    timestamp_end=segment.end,
                    text=text
                ))
        
        if rows:
            emb = LocalEmbeddingProvider()
            vecs = emb.embed([r.text for r in rows])
            ensure_collection(len(vecs[0]))
            upsert([PointStruct(id=str(uuid.uuid4()), vector=v, payload={
                "course_id": r.course_id, "source_id": r.source_id, "source_name": src.name,
                "timestamp_start": r.timestamp_start, "timestamp_end": r.timestamp_end,
                "text": r.text}) for r, v in zip(rows, vecs)])
            db.add_all(rows)
            
        src.status = "COMPLETED"
    except Exception as e:
        log.exception("youtube ingest failed")
        src.status = "FAILED"; src.error = "YouTube processing failed."
    db.commit(); db.close()

def ingest_pptx(source_id: str):
    db = SessionLocal()
    src = db.get(Source, source_id)
    try:
        src.status = "PROCESSING"
        db.commit()
        from pptx import Presentation
        prs = Presentation(src.path)
        src.pages = len(prs.slides)
        rows = []
        for s_idx, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    texts.append(f"Speaker notes: {notes}")
            slide_text = "\n".join(texts).strip()
            for t in chunk_text(slide_text):
                if len(t.strip()) > 20:
                    rows.append(Chunk(
                        source_id=src.id,
                        course_id=src.course_id,
                        slide=s_idx,
                        page=s_idx,
                        text=t
                    ))
        if rows:
            emb = LocalEmbeddingProvider()
            vecs = emb.embed([r.text for r in rows])
            ensure_collection(len(vecs[0]))
            upsert([PointStruct(id=str(uuid.uuid4()), vector=v, payload={
                "course_id": r.course_id, "source_id": r.source_id, "source_name": src.name,
                "slide": r.slide, "page": r.page, "text": r.text}) for r, v in zip(rows, vecs)])
            db.add_all(rows)
        src.status = "COMPLETED"
    except Exception as e:
        log.exception("pptx ingest failed")
        src.status = "FAILED"
        src.error = "PPTX processing failed."
    finally:
        db.commit()
        db.close()

