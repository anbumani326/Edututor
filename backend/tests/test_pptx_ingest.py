from pptx import Presentation
from app.core import SessionLocal
from app.models import User, Course, Source, Chunk
from app.ingestion import ingest_pptx

def test_pptx_ingestion(tmp_path, monkeypatch):
    # Mock embedding provider to prevent heavy CPU model load in unit test
    class MockEmb:
        def embed(self, texts):
            return [[0.1] * 384 for _ in texts]
    monkeypatch.setattr("app.ingestion.LocalEmbeddingProvider", MockEmb)
    monkeypatch.setattr("app.ingestion.ensure_collection", lambda dim: None)
    monkeypatch.setattr("app.ingestion.upsert", lambda points: None)

    # Create test PPTX
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Introduction to Neural Networks"
    slide1.placeholders[1].text = "Perceptrons and multi-layer feedforward architectures."
    
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Backpropagation Algorithm"
    slide2.placeholders[1].text = "Gradient descent calculation with chain rule of calculus."
    
    pptx_file = str(tmp_path / "test_presentation.pptx")
    prs.save(pptx_file)
    
    import uuid
    db = SessionLocal()
    email = f"pptx_tester_{uuid.uuid4().hex[:8]}@example.com"
    u = User(email=email, password_hash="hash")
    db.add(u)
    db.commit()
    
    c = Course(owner_id=u.id, name="Test PPTX Course")
    db.add(c)
    db.commit()
    
    src = Source(course_id=c.id, name="test_presentation.pptx", source_type="pptx", path=pptx_file, status="PENDING")
    db.add(src)
    db.commit()
    
    # Ingest
    ingest_pptx(src.id)
    
    db.refresh(src)
    assert src.status == "COMPLETED"
    assert src.pages == 2
    
    chunks = db.query(Chunk).filter_by(source_id=src.id).all()
    assert len(chunks) >= 2
    slides = [ch.slide for ch in chunks]
    assert 1 in slides and 2 in slides
